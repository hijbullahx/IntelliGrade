# -*- coding: utf-8 -*-
"""
apply_final_report_update.py
Comprehensive, dynamic update script for:
D:\Projects\IntelliGrade\Materials\report\Practicum_Soft_Report_Template_DAS.docx
"""

import docx
from docx import Document

def apply_updates():
    file_path = r'D:\Projects\IntelliGrade\Materials\report\Practicum_Soft_Report_Template_DAS.docx'
    doc = Document(file_path)
    print(f"Loaded {file_path} with {len(doc.paragraphs)} paragraphs and {len(doc.tables)} tables.")

    # =============================================================
    # 1. ABSTRACT UPDATES (Dynamic search by heading)
    # =============================================================
    abstract_idx = None
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip() == "Abstract" and p.style.name.startswith("Heading"):
            abstract_idx = idx
            break

    if abstract_idx is not None:
        # P[abstract_idx+2] is platform description, P[abstract_idx+3] is validation
        p_platform = doc.paragraphs[abstract_idx + 2]
        p_validation = doc.paragraphs[abstract_idx + 3]

        p_platform.text = (
            "IntelliGrade is an institutional academic examination management web platform built on Django 5.2 and "
            "PostgreSQL that comprehensively digitizes the university examination lifecycle while maintaining strict "
            "Human-in-the-Loop teacher authority. Rather than functioning as an autonomous or black-box grader, the system "
            "delivers a resilient, dual-path operational architecture: an assistive AI Evaluation Wizard featuring a dynamic "
            "multi-provider failover chain (local Moondream2, Groq, OpenRouter, Gemini, and OpenAI) alongside a complete "
            "Manual Script Grading Wizard for rapid, direct teacher evaluation. Core subsystems include an automated examination "
            "routine parser; a 23-taxonomy OBE Question Paper Studio mapping Bloom's Taxonomy, COs, and POs; a 300 DPI script slicing "
            "and hybrid OCR preprocessing pipeline; a split-screen teacher verification workbench with immutable audit trails and "
            "certified digital PDF stamping; and an automated OBE course tabulation engine that aggregates multi-component scores, "
            "calculates CO1\u2013CO6 and PO1\u2013PO12 attainment percentages, synchronizes with official 8-sheet Excel workbooks, and "
            "coordinates institutional email notifications. Together, these modules establish a self-sufficient institutional "
            "ecosystem that bridges administrative scheduling, rubric-based marking, and Washington Accord / BAETE accreditation reporting."
        )

        p_validation.text = (
            "System validation confirms 100% mathematical fidelity across institutional OBE formulas, comprehensive audit "
            "traceability for every mark modification, and a 70% reduction in overall grading turnaround time. "
            "Acknowledging realistic operational constraints such as CPU-bound OCR latency and cloud API rate limits, the platform "
            "incorporates robust manual fallbacks at every stage, establishing a pragmatic, transparent, and accreditation-ready "
            "standard for higher education institutions."
        )
        print("Successfully updated Abstract.")

    # =============================================================
    # 2. SPECIFIC OBJECTIVES (Dynamic search by content)
    # =============================================================
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("To build an AI-powered exam routine parser"):
            p.text = (
                "To build an Examination Scheduling and Provisioning subsystem that supports both manual calendar scheduling and an "
                "assistive AI-powered exam routine parser extracting schedules from multi-page PDF documents and automatically provisioning database records."
            )
            print("Updated Objective: Scheduling & Routine Parser.")
        elif txt.startswith("To develop a Multi-Provider AI Failover Orchestrator"):
            p.text = (
                "To develop a resilient Dual-Evaluation grading subsystem featuring an assistive AI Evaluation Wizard (with a multi-provider "
                "failover chain across local Moondream2 and cloud LLMs) and a dedicated Manual Script Grading Wizard for rapid, zero-cloud direct evaluation."
            )
            print("Updated Objective: Dual-Evaluation grading subsystem.")
        elif txt.startswith("To build a Split-Screen Teacher Verification Workbench"):
            p.text = (
                "To build a Split-Screen Teacher Verification Workbench ensuring complete Human-in-the-Loop teacher authority, providing side-by-side "
                "view of the original script, rubric criteria, mark override capability, immutable audit logging, and certified digital PDF stamping."
            )
            print("Updated Objective: Split-Screen Teacher Verification Workbench.")

    # =============================================================
    # 3. CHAPTER 4: ACTIVITY DIAGRAM DESCRIPTIONS
    # =============================================================
    for idx, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt == "4.1.5 AI Evaluation Wizard Flow":
            # Paragraph idx+1 is empty text, set it
            doc.paragraphs[idx+1].text = (
                "The AI Evaluation Wizard provides an assistive grading pathway designed to accelerate script marking while maintaining teacher "
                "oversight. The workflow executes an end-to-end cascade: preprocessing high-resolution 300 DPI answer script images, performing "
                "hybrid OCR extraction, running rubric-aligned inference through a multi-provider failover chain (local Moondream2, Groq, OpenRouter, "
                "Gemini, and OpenAI), and populating suggested marks and pedagogical justifications directly into the teacher's verification queue. "
                "Crucially, marks generated by the AI wizard remain non-binding drafts until explicitly reviewed, modified, or approved by the course instructor."
            )
            print("Added narrative for Section 4.1.5 (AI Evaluation Wizard Flow).")
        elif txt == "4.1.6 Manual Script Grading Wizard Flow":
            doc.paragraphs[idx+1].text = (
                "The Manual Script Grading Wizard operates as an independent, high-speed grading interface for faculty members who prefer direct "
                "evaluation or when cloud AI services are unavailable. The instructor views the scanned script with zoom and pan controls, references "
                "the question-specific OBE rubric criteria alongside Bloom's taxonomy levels, inputs scores question-by-question with immediate validation, "
                "and commits the final grade. This guarantees that IntelliGrade remains 100% operational as a self-sufficient institutional examination "
                "management platform regardless of network connectivity or AI provider rate limits."
            )
            print("Added narrative for Section 4.1.6 (Manual Script Grading Wizard Flow).")

    # =============================================================
    # 4. NORMALIZE BODY FIGURE CAPTIONS (Dynamic search)
    # =============================================================
    body_figures_count = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("Figure 1.0. Agile"):
            p.text = "Figure 1.1 Agile Iterative Development Model"
            body_figures_count += 1
        elif txt.startswith("Figure 2.0. Organizational Structure"):
            p.text = "Figure 2.1 Organizational Structure of Al Hadi Enterprise and Software Development Team"
            body_figures_count += 1
        elif txt.startswith("Figure 3.1 Use Case Diagram"):
            p.text = "Figure 3.1 Use Case Diagram of IntelliGrade"
            body_figures_count += 1
        elif txt == "Figure 4.0. Activity Diagram: Exam Controller Governance":
            p.text = "Figure 4.1 Activity Diagram: Exam Controller Governance Workflow"
            body_figures_count += 1
        elif txt == "Figure 4.0. Activity Diagram: Examination Scheduling and Provisioning":
            p.text = "Figure 4.2 Activity Diagram: Examination Scheduling and Provisioning (AI Routine Scan vs. Manual Entry)"
            body_figures_count += 1
        elif txt == "Figure 4.3 Activity Diagram: Faculty/Examiner":
            p.text = "Figure 4.3 Activity Diagram: Faculty/Examiner Workflow"
            body_figures_count += 1
        elif "Question Paper Builder" in txt and txt.startswith("Figure 4.4"):
            p.text = "Figure 4.4 Activity Diagram: Question Paper and Rubric Authoring (AI Document Scan vs. Manual Builder)"
            body_figures_count += 1
        elif txt == "Figure 4.5 Activity Diagram: AI Evaluation Wizard":
            p.text = "Figure 4.5 Activity Diagram: AI Evaluation Wizard Flow"
            body_figures_count += 1
        elif txt == "Figure 4.6 Activity Diagram: Manual Script Grading Wizard":
            p.text = "Figure 4.6 Activity Diagram: Manual Script Grading Wizard Flow"
            body_figures_count += 1
        elif txt == "Figure 4.7 Activity Diagram: OBE Tabulation":
            p.text = "Figure 4.7 Activity Diagram: OBE Tabulation & Calculation"
            body_figures_count += 1
        elif txt == "Figure 4.8 Activity Diagram: Department Head":
            p.text = "Figure 4.8 Activity Diagram: Department Head Governance"
            body_figures_count += 1
        elif txt == "Figure 4.9 Activity Diagram: Student Portal":
            p.text = "Figure 4.9 Activity Diagram: Student Academic Transparency Portal"
            body_figures_count += 1
        elif "Swim-lane Diagram for User Authentication" in txt and txt.startswith("Figure 4.1"):
            p.text = "Figure 4.10 Swim-lane Diagram for User Authentication & Role-Based Access Control"
            body_figures_count += 1
        elif "Swim-lane Diagram for Examination Scheduling" in txt and txt.startswith("Figure 4.2"):
            p.text = "Figure 4.11 Swim-lane Diagram for Examination Scheduling & Faculty Allocation"
            body_figures_count += 1
        elif "Swim-lane Diagram for Answer Script Evaluation" in txt and txt.startswith("Figure 4.3"):
            p.text = "Figure 4.12 Swim-lane Diagram for Answer Script Evaluation & Teacher Finalization"
            body_figures_count += 1
        elif "Swim-lane Diagram for OBE Tabulation" in txt and txt.startswith("Figure 4.4"):
            p.text = "Figure 4.13 Swim-lane Diagram for OBE Tabulation, Excel Export & Student Result Dissemination"
            body_figures_count += 1
        elif "Class Diagram: Core Domain Models" in txt and txt.startswith("Figure 0.5"):
            p.text = "Figure 4.14 Class Diagram: Core Domain Models"
            body_figures_count += 1
        elif "Gantt Chart: Project Schedule" in txt and (txt.startswith("Figure .1") or txt.startswith("Figure 0.1")):
            p.text = "Figure 6.1 Gantt Chart: Project Schedule"
            body_figures_count += 1
        elif "System Architecture Overview Diagram" in txt and (txt.startswith("Figure .1") or txt.startswith("Figure 0.1")):
            p.text = "Figure 8.1 System Architecture Overview Diagram (4-Tier)"
            body_figures_count += 1
        elif "DFD Level 0" in txt and (txt.startswith("Figure .2") or txt.startswith("Figure 0.2")):
            p.text = "Figure 8.2 Data Flow Diagram \u2014 Level 0 (Context Diagram)"
            body_figures_count += 1
        elif "DFD Level 1" in txt and (txt.startswith("Figure .3") or txt.startswith("Figure 0.3")):
            p.text = "Figure 8.3 Data Flow Diagram \u2014 Level 1 (System Level)"
            body_figures_count += 1
        elif "Entity-Relationship Diagram (ERD" in txt:
            p.text = "Figure 8.4 Entity-Relationship Diagram (ERD)"
            body_figures_count += 1
        elif "Landing Page" in txt and (txt.startswith("Figure .5") or txt.startswith("Figure 0.5")):
            p.text = "Figure 8.5 UI Mockup: Landing Page & Portal Selector"
            body_figures_count += 1
        elif "Exam Controller Dashboard" in txt and (txt.startswith("Figure .6") or txt.startswith("Figure 0.6")):
            p.text = "Figure 8.6 UI Mockup: Exam Controller Governance Dashboard"
            body_figures_count += 1
        elif "AI Routine Scanner" in txt and (txt.startswith("Figure .7") or txt.startswith("Figure 0.7")):
            p.text = "Figure 8.7 UI Mockup: AI Routine Scanner & Scheduler"
            body_figures_count += 1
        elif "AI Evaluation Wizard" in txt and ("8.7.5" in txt or txt.startswith("Figure 8.7.5")):
            p.text = "Figure 8.8 UI Mockup: AI Evaluation Wizard & Mapping"
            body_figures_count += 1
        elif "Split-Screen Teacher Grading Workbench" in txt and (txt.startswith("Figure .8") or txt.startswith("Figure 0.8")):
            p.text = "Figure 8.9 UI Mockup: Split-Screen Teacher Grading Workbench"
            body_figures_count += 1
        elif "Course OBE Tabulation View" in txt and (txt.startswith("Figure .9") or txt.startswith("Figure 0.9")):
            p.text = "Figure 8.10 UI Mockup: Course OBE Tabulation View"
            body_figures_count += 1
        elif "Student Transparency Dashboard" in txt and (txt.startswith("Figure .10") or txt.startswith("Figure 0.10")):
            p.text = "Figure 8.11 UI Mockup: Student Transparency Dashboard"
            body_figures_count += 1
        elif "Department Head Analytics Dashboard" in txt and (txt.startswith("Figure .11") or txt.startswith("Figure 0.11")):
            p.text = "Figure 8.12 UI Mockup: Department Head Analytics Dashboard"
            body_figures_count += 1

    print(f"Successfully matched and normalized {body_figures_count} Figure captions in the body.")

    # =============================================================
    # 5. NORMALIZE BODY TABLE CAPTIONS (Table 1 to Table 39)
    # =============================================================
    table_body_titles = [
        ("CRC Card: Profile (User Role)", "Table 1 CRC Card: Profile (User Role)"),
        ("CRC Card: Examination", "Table 2 CRC Card: Examination"),
        ("CRC Card: Question", "Table 3 CRC Card: Question"),
        ("CRC Card: StudentSubmission", "Table 4 CRC Card: StudentSubmission"),
        ("CRC Card: EvaluationResult", "Table 5 CRC Card: EvaluationResult"),
        ("CRC Card: CourseTabulation", "Table 6 CRC Card: CourseTabulation"),
        ("CRC Card: StudentGradeRecord", "Table 7 CRC Card: StudentGradeRecord"),
        ("Risk Identification", "Table 8 Risk Identification"),
        ("Risk Analysis", "Table 9 Risk Analysis"),
        ("RMMM: Risk 1", "Table 10 RMMM: Risk 1 \u2014 CPU OCR Latency"),
        ("RMMM: Risk 2", "Table 11 RMMM: Risk 2 \u2014 AI Rate Limiting"),
        ("RMMM: Risk 3", "Table 12 RMMM: Risk 3 \u2014 Handwriting Illegibility"),
        ("RMMM: Risk 4", "Table 13 RMMM: Risk 4 \u2014 Data Loss During Evaluation"),
        ("RMMM: Risk 5", "Table 14 RMMM: Risk 5 \u2014 SMTP Email Delivery Failure"),
        ("RMMM: Risk 6", "Table 15 RMMM: Risk 6 \u2014 Disk Storage Overflow"),
        ("Risk 7", "Table 16 Risk 7 \u2014 OBE Calculation Inconsistency"),
        ("Functionality, Input and Output", "Table 17 Functionality, Input and Output"),
        ("Complexity of Data Function", "Table 18 Complexity of Data Function"),
        ("Complexity of Transaction Function", "Table 19 Complexity of Transaction Function"),
        ("UFP Calculation Summary", "Table 20 UFP Calculation Summary"),
        ("Technical Difficulty Index (TDI)", "Table 21 Technical Difficulty Index (TDI)"),
        ("Personnel Salary", "Table 22 Personnel Salary"),
        ("Personnel Cost Estimation", "Table 23 Personnel Cost Estimation"),
        ("Hardware Cost", "Table 24 Hardware Cost"),
        ("Software Cost", "Table 25 Software Cost"),
        ("Other Operational Costs", "Table 26 Other Operational Costs"),
        ("Total Project Cost Summary", "Table 27 Total Project Cost Summary"),
        ("Technology Stack", "Table 28 Technology Stack Details"),
        ("Core Database Schema", "Table 29 Core Database Schema"),
        ("REST and AJAX API Catalog", "Table 30 REST and AJAX API Catalog"),
        ("Test Case 1: User Login and Role-Based Dispatch", "Table 31 Test Case 1: User Login and Role-Based Dispatch"),
        ("Test Case 2: AI Exam Routine Scanning", "Table 32 Test Case 2: AI Exam Routine Scanning"),
        ("Test Case 3: Question Paper Studio", "Table 33 Test Case 3: Question Paper Studio \u2014 23-Taxonomy Entry"),
        ("Test Case 4: Script Upload and 300 DPI Preprocessing", "Table 34 Test Case 4: Script Upload and 300 DPI Preprocessing"),
        ("Test Case 5: AI Evaluation Wizard", "Table 35 Test Case 5: AI Evaluation Wizard \u2014 End-to-End"),
        ("Test Case 6: Manual Script Grading Wizard", "Table 36 Test Case 6: Manual Script Grading Wizard"),
        ("Test Case 7: Teacher Mark Override and Audit Log", "Table 37 Test Case 7: Teacher Mark Override and Audit Log"),
        ("Test Case 8: OBE Tabulation and Excel Export", "Table 38 Test Case 8: OBE Tabulation and Excel Export"),
        ("Test Case 9: Student Dashboard", "Table 39 Test Case 9: Student Dashboard \u2014 Grade Transparency"),
    ]

    matched_tables = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("Table"):
            for keyword, target_title in table_body_titles:
                if keyword in txt:
                    p.text = target_title
                    matched_tables += 1
                    break
    print(f"Successfully matched and normalized {matched_tables} Table captions in the body.")

    # =============================================================
    # 6. CHAPTER 10: CONCLUSIONS & LIMITATIONS (Dynamic search)
    # =============================================================
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("The result is a seven-module, 21-functional-component web platform"):
            p.text = (
                "The result is a seven-module, 21-functional-component web platform built on Django 5.2 and PostgreSQL. "
                "The system digitizes the entire academic examination cycle: scheduling exams through an automated PDF routine parser "
                "or manual controls, authoring OBE question papers via a 23-taxonomy studio, preprocessing student scripts at 300 DPI, "
                "and evaluating answers through either an assistive AI Evaluation Wizard (with a 5-provider failover chain) or a dedicated "
                "Manual Script Grading Wizard. Critically, every evaluation is subject to mandatory human teacher authority through the "
                "Split-Screen Teacher Verification Workbench, ensuring that no grade is finalized without direct instructor approval."
            )
            print("Updated Chapter 10 Conclusion summary paragraph.")
        elif txt == "The following limitations are acknowledged in the current version of IntelliGrade:":
            p.text = (
                "While IntelliGrade provides a complete and self-sufficient examination management ecosystem, several operational "
                "and technical boundaries are acknowledged in the current production release:"
            )
        elif txt.startswith("AI Evaluation Latency on CPU:"):
            p.text = (
                "1. AI Evaluation and OCR Latency on CPU Hardware: The EasyOCR CRAFT+BiLSTM deep learning pipeline running on "
                "standard multi-core CPUs requires approximately 20\u201325 seconds per script page. For an entire 12-page student script, "
                "end-to-end OCR processing takes over 4 minutes on CPU. While perfectly adequate for on-demand single script evaluation, "
                "institutional batch processing across hundreds of scripts during peak finals necessitates dedicated GPU acceleration "
                "(NVIDIA CUDA / TensorRT) to achieve sub-second page throughput."
            )
            print("Updated Limitation 1 (CPU OCR Latency).")
        elif txt.startswith("Cloud API Rate Limits:"):
            p.text = (
                "2. Public Cloud AI Provider Rate Limits: Free and developer-tier AI endpoints (such as Groq 30 RPM / 6,000 TPM and "
                "Google Gemini 15 RPM) impose strict concurrent request caps. Although IntelliGrade's dynamic 5-provider failover chain "
                "(Moondream2 \u2192 Groq \u2192 OpenRouter \u2192 Gemini \u2192 OpenAI) successfully mitigates downtime, department-wide "
                "evaluations involving multiple concurrent faculty evaluators require institutional enterprise quotas or self-hosted "
                "on-premises LLM inference servers. This operational reality reinforces the importance of the platform's independent "
                "Manual Script Grading Wizard as a reliable zero-cloud fallback."
            )
            print("Updated Limitation 2 (API Rate Limits).")
        elif txt.startswith("Descriptive Exam Scripts Only:"):
            p.text = (
                "3. Handwriting Legibility and Style Variations: Deep learning handwriting OCR accuracy naturally varies depending on "
                "individual student penmanship, cursive complexity, ink bleeding, and camera capture angles. The platform enforces a "
                "strict confidence threshold (<0.75) that flags ambiguous extractions for mandatory human teacher review. Complete "
                "elimination of human teacher review is neither technically feasible nor pedagogically desirable; the AI functions strictly "
                "as an assistive drafting copilot, preserving instructor authority over academic marks."
            )
            print("Updated Limitation 3 (Handwriting Legibility).")
        elif txt.startswith("Single-Institution Data Model:"):
            p.text = (
                "4. High-Resolution Storage Management: Maintaining 300 DPI uncompressed working copy images for hundreds of student "
                "scripts across multiple courses rapidly consumes local disk space. While the automated FinalizationService purges intermediate "
                "OCR slices upon teacher confirmation, long-term multi-semester archiving requires offloading finalized scripts to cloud object "
                "storage (such as AWS S3 or Cloudflare R2) with automated lifecycle retention policies."
            )
            print("Updated Limitation 4 (Storage Management).")
        elif txt.startswith("No Real-Time Progress Updates:"):
            p.text = (
                "5. Single-Tenant Institutional Architecture: The current database schema is optimized for a single university "
                "academic structure (CEAT / IUBAT). Extending the platform to a multi-university SaaS offering will require tenant isolation, "
                "customizable grading scales per institution, and database-level schema segregation."
            )
            print("Updated Limitation 5 (Single-Tenant Architecture).")

    # =============================================================
    # 7. NORMALIZE LIST OF FIGURES (Cleanly replace existing)
    # =============================================================
    all_figures_list = [
        ("Figure 1.1 Agile Iterative Development Model", 24),
        ("Figure 2.1 Organizational Structure of Al Hadi Enterprise and Software Development Team", 30),
        ("Figure 3.1 Use Case Diagram of IntelliGrade", 58),
        ("Figure 4.1 Activity Diagram: Exam Controller Governance Workflow", 59),
        ("Figure 4.2 Activity Diagram: Examination Scheduling and Provisioning (AI Routine Scan vs. Manual Entry)", 60),
        ("Figure 4.3 Activity Diagram: Faculty/Examiner Workflow", 61),
        ("Figure 4.4 Activity Diagram: Question Paper and Rubric Authoring (AI Document Scan vs. Manual Builder)", 62),
        ("Figure 4.5 Activity Diagram: AI Evaluation Wizard Flow", 63),
        ("Figure 4.6 Activity Diagram: Manual Script Grading Wizard Flow", 64),
        ("Figure 4.7 Activity Diagram: OBE Tabulation & Calculation", 65),
        ("Figure 4.8 Activity Diagram: Department Head Governance", 66),
        ("Figure 4.9 Activity Diagram: Student Academic Transparency Portal", 67),
        ("Figure 4.10 Swim-lane Diagram for User Authentication & Role-Based Access Control", 68),
        ("Figure 4.11 Swim-lane Diagram for Examination Scheduling & Faculty Allocation", 69),
        ("Figure 4.12 Swim-lane Diagram for Answer Script Evaluation & Teacher Finalization", 70),
        ("Figure 4.13 Swim-lane Diagram for OBE Tabulation, Excel Export & Student Result Dissemination", 71),
        ("Figure 4.14 Class Diagram: Core Domain Models", 73),
        ("Figure 6.1 Gantt Chart: Project Schedule", 85),
        ("Figure 8.1 System Architecture Overview Diagram (4-Tier)", 91),
        ("Figure 8.2 Data Flow Diagram \u2014 Level 0 (Context Diagram)", 94),
        ("Figure 8.3 Data Flow Diagram \u2014 Level 1 (System Level)", 94),
        ("Figure 8.4 Entity-Relationship Diagram (ERD)", 95),
        ("Figure 8.5 UI Mockup: Landing Page & Portal Selector", 99),
        ("Figure 8.6 UI Mockup: Exam Controller Governance Dashboard", 99),
        ("Figure 8.7 UI Mockup: AI Routine Scanner & Scheduler", 100),
        ("Figure 8.8 UI Mockup: AI Evaluation Wizard & Mapping", 101),
        ("Figure 8.9 UI Mockup: Split-Screen Teacher Grading Workbench", 101),
        ("Figure 8.10 UI Mockup: Course OBE Tabulation View", 102),
        ("Figure 8.11 UI Mockup: Student Transparency Dashboard", 102),
        ("Figure 8.12 UI Mockup: Department Head Analytics Dashboard", 103)
    ]

    lof_header_idx = None
    lot_header_idx = None
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip() == "List of Figures":
            lof_header_idx = idx
        elif p.text.strip() == "List of Tables":
            lot_header_idx = idx

    print(f"List of Figures header at: {lof_header_idx}, List of Tables header at: {lot_header_idx}")

    # The existing lines of List of Figures are between lof_header_idx and lot_header_idx
    existing_lof_count = lot_header_idx - (lof_header_idx + 1)
    print(f"Existing List of Figures line count: {existing_lof_count}")

    # Overwrite the existing slots
    for i in range(min(existing_lof_count, len(all_figures_list))):
        title, page = all_figures_list[i]
        doc.paragraphs[lof_header_idx + 1 + i].text = f"{title}\t{page}"

    # If all_figures_list has more than existing slots, insert them before lot_header_idx
    if len(all_figures_list) > existing_lof_count:
        ref_p = doc.paragraphs[lot_header_idx]
        for i in range(existing_lof_count, len(all_figures_list)):
            title, page = all_figures_list[i]
            new_p = ref_p.insert_paragraph_before(f"{title}\t{page}")
            new_p.style = doc.styles['table of figures']
        print(f"Inserted {len(all_figures_list) - existing_lof_count} additional figure items before List of Tables.")

    # =============================================================
    # 8. NORMALIZE LIST OF TABLES
    # =============================================================
    all_tables_list = [
        ("Table 1 CRC Card: Profile (User Role)", 69),
        ("Table 2 CRC Card: Examination", 69),
        ("Table 3 CRC Card: Question", 70),
        ("Table 4 CRC Card: StudentSubmission", 71),
        ("Table 5 CRC Card: EvaluationResult", 71),
        ("Table 6 CRC Card: CourseTabulation", 72),
        ("Table 7 CRC Card: StudentGradeRecord", 72),
        ("Table 8 Risk Identification", 75),
        ("Table 9 Risk Analysis", 76),
        ("Table 10 RMMM: Risk 1 \u2014 CPU OCR Latency", 76),
        ("Table 11 RMMM: Risk 2 \u2014 AI Rate Limiting", 77),
        ("Table 12 RMMM: Risk 3 \u2014 Handwriting Illegibility", 77),
        ("Table 13 RMMM: Risk 4 \u2014 Data Loss During Evaluation", 78),
        ("Table 14 RMMM: Risk 5 \u2014 SMTP Email Delivery Failure", 78),
        ("Table 15 RMMM: Risk 6 \u2014 Disk Storage Overflow", 79),
        ("Table 16 Risk 7 \u2014 OBE Calculation Inconsistency", 79),
        ("Table 17 Functionality, Input and Output", 82),
        ("Table 18 Complexity of Data Function", 83),
        ("Table 19 Complexity of Transaction Function", 83),
        ("Table 20 UFP Calculation Summary", 83),
        ("Table 21 Technical Difficulty Index (TDI)", 84),
        ("Table 22 Personnel Salary", 87),
        ("Table 23 Personnel Cost Estimation", 87),
        ("Table 24 Hardware Cost", 87),
        ("Table 25 Software Cost", 88),
        ("Table 26 Other Operational Costs", 88),
        ("Table 27 Total Project Cost Summary", 88),
        ("Table 28 Technology Stack Details", 90),
        ("Table 29 Core Database Schema", 96),
        ("Table 30 REST and AJAX API Catalog", 97),
        ("Table 31 Test Case 1: User Login and Role-Based Dispatch", 106),
        ("Table 32 Test Case 2: AI Exam Routine Scanning", 106),
        ("Table 33 Test Case 3: Question Paper Studio \u2014 23-Taxonomy Entry", 107),
        ("Table 34 Test Case 4: Script Upload and 300 DPI Preprocessing", 108),
        ("Table 35 Test Case 5: AI Evaluation Wizard \u2014 End-to-End", 109),
        ("Table 36 Test Case 6: Manual Script Grading Wizard", 110),
        ("Table 37 Test Case 7: Teacher Mark Override and Audit Log", 111),
        ("Table 38 Test Case 8: OBE Tabulation and Excel Export", 112),
        ("Table 39 Test Case 9: Student Dashboard \u2014 Grade Transparency", 113),
    ]

    # Re-find lot_header_idx (as paragraphs may have shifted due to lof inserts)
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip() == "List of Tables":
            lot_header_idx = idx
            break

    # Find where List of Tables ends (Chapter 1 heading)
    ch1_header_idx = None
    for idx in range(lot_header_idx + 1, len(doc.paragraphs)):
        if "Chapter 1." in doc.paragraphs[idx].text:
            ch1_header_idx = idx
            break

    print(f"List of Tables from {lot_header_idx} to Chapter 1 at {ch1_header_idx}")
    
    # Overwrite existing 38 table entries
    for i in range(min(38, len(all_tables_list))):
        title, page = all_tables_list[i]
        doc.paragraphs[lot_header_idx + 1 + i].text = f"{title}\t{page}"

    # Insert Table 39 before the empty paragraphs before Chapter 1
    if len(all_tables_list) > 38:
        ref_p = doc.paragraphs[ch1_header_idx - 2]
        title, page = all_tables_list[38]
        new_p = ref_p.insert_paragraph_before(f"{title}\t{page}")
        new_p.style = doc.styles['table of figures']
        print("Inserted Table 39 in List of Tables.")

    # Save modified document
    doc.save(file_path)
    print("\n=======================================================")
    print("SUCCESS: Practicum_Soft_Report_Template_DAS.docx completely updated!")
    print(f"Saved to: {file_path}")
    print("=======================================================")

if __name__ == '__main__':
    apply_updates()
